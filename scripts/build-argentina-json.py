
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path.cwd()
EXCEL_FILE = ROOT / "local.xlsx"
OUTPUT_FILE = ROOT / "argentina.json"
SHEET_NAME = "Etiquetas_localidades"

COLUMN_MAP = {
    "código de jurisdicción": "jurisdiction_code",
    "jurisdicción": "jurisdiction",
    "código de departamento": "department_code",
    "departamento": "department",
    "código de aglomerado": "agglomeration_code",
    "aglomerado": "agglomeration",
    "código de localidad": "locality_id",
    "localidad": "locality",
    "tipo de localidad": "locality_type",
    "código de gobierno local": "local_government_code",
    "gobierno local": "local_government",
    "categoría de gobierno local": "local_government_category",
}


def normalize_header(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def get_value(cell, is_code=False):
    if cell is None or cell.value is None:
        return None

    value = cell.value

    if isinstance(value, str):
        value = value.strip()
        return value or None

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, (int, float)):
        if isinstance(value, float):
            if not math.isfinite(value):
                return None
            if value.is_integer():
                value = int(value)

        if is_code and isinstance(value, int):
            number_format = cell.number_format or ""
            if re.fullmatch(r"0+", number_format):
                return str(value).zfill(len(number_format))

        return str(value)

    value = str(value).strip()
    return value or None


def main():
    if not EXCEL_FILE.is_file():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    print(f"Reading workbook: {EXCEL_FILE}")

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{SHEET_NAME}' not found. "
                f"Available worksheets: {workbook.sheetnames}"
            )

        worksheet = workbook[SHEET_NAME]
        iterator = worksheet.iter_rows()
        header_row = next(iterator, None)

        if not header_row:
            raise ValueError("The worksheet is empty.")

        header_indexes = {}

        for index, cell in enumerate(header_row):
            header = normalize_header(cell.value)
            if header:
                header_indexes[header] = index

        missing = [
            column for column in COLUMN_MAP
            if column not in header_indexes
        ]

        if missing:
            raise ValueError(
                "Missing required columns: "
                + ", ".join(missing)
            )

        records = []
        seen_ids = set()
        skipped_rows = 0

        for row_number, row in enumerate(iterator, start=2):
            record = {}

            for source_column, output_column in COLUMN_MAP.items():
                index = header_indexes[source_column]
                cell = row[index] if index < len(row) else None

                record[output_column] = get_value(
                    cell,
                    is_code=("código" in source_column),
                )

            locality_id = record["locality_id"]
            locality_name = record["locality"]

            if not locality_id or not locality_name:
                skipped_rows += 1
                continue

            record["locality_id"] = str(locality_id)

            if record["locality_id"] in seen_ids:
                raise ValueError(
                    f"Duplicate locality ID '{record['locality_id']}' "
                    f"at Excel row {row_number}."
                )

            seen_ids.add(record["locality_id"])
            records.append(record)

    finally:
        workbook.close()

    if not records:
        raise ValueError("No valid locality records were found.")

    records.sort(
        key=lambda item: (
            item["jurisdiction_code"] or "",
            item["department_code"] or "",
            item["locality_id"],
        )
    )

    # Write the same JSON structure without indentation or extra spaces.
    temporary_file = OUTPUT_FILE.with_suffix(".json.tmp")

    with temporary_file.open(
        "w",
        encoding="utf-8",
        newline="\n",
    ) as file:
        json.dump(
            records,
            file,
            ensure_ascii=False,
            separators=(",", ":"),
            allow_nan=False,
        )
        file.write("\n")

    # Verify that the generated file is valid JSON.
    with temporary_file.open("r", encoding="utf-8") as file:
        validated_records = json.load(file)

    if len(validated_records) != len(records):
        raise ValueError("JSON validation failed.")

    temporary_file.replace(OUTPUT_FILE)

    print("SUCCESS: argentina.json generated and validated.")
    print(f"Valid localities: {len(records)}")
    print(f"Unique locality IDs: {len(seen_ids)}")
    print(f"Skipped blank/footer rows: {skipped_rows}")
    print(f"Final file size: {OUTPUT_FILE.stat().st_size:,} bytes")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        sys.exit(1)
